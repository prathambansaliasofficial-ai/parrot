"""Extraction and passage indexing for educational documents."""

from __future__ import annotations

import hashlib
import io
import logging
import re
import threading
import zipfile
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import PurePosixPath
from uuid import UUID

import pymupdf as fitz
from docx import Document
from fastapi import HTTPException, status

from app.core.config import get_settings
from app.core.supabase import get_supabase_client

logger = logging.getLogger(__name__)
_locks_guard = threading.Lock()
_resource_locks: dict[str, threading.Lock] = {}
_CHUNK_CHARS = 1200


@dataclass
class ExtractedSection:
    text: str
    page_number: int | None
    section_ref: str | None
    topic: str | None
    subtopic: str | None = None


def _lock_for(resource_id: str) -> threading.Lock:
    with _locks_guard:
        return _resource_locks.setdefault(resource_id, threading.Lock())


def _clean_text(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _looks_like_heading(text: str) -> bool:
    value = text.strip()
    return bool(value and len(value) <= 100 and (value.isupper() or re.match(r"^(chapter|unit|module|section)\s+[\w.-]+", value, re.I)))


def extract_document(content: bytes, filename: str, content_type: str) -> tuple[list[ExtractedSection], int]:
    """Return structure-preserving sections and page count; reject empty/scanned files."""
    settings = get_settings()
    extension = PurePosixPath(filename).suffix.lower()
    sections: list[ExtractedSection] = []
    page_count = 1
    if extension == ".pdf" and content_type == "application/pdf":
        try:
            document = fitz.open(stream=content, filetype="pdf")
        except Exception as exc:
            raise ValueError("The PDF is corrupted or unreadable.") from exc
        with document:
            if document.page_count > settings.max_document_pages:
                raise ValueError(f"Document exceeds the {settings.max_document_pages}-page limit.")
            page_count = document.page_count
            topic: str | None = None
            char_count = 0
            for page_index, page in enumerate(document):
                lines = [line.strip() for line in page.get_text("text").splitlines() if line.strip()]
                body: list[str] = []
                for line in lines:
                    if _looks_like_heading(line):
                        if body:
                            text = _clean_text("\n".join(body))
                            if text:
                                sections.append(ExtractedSection(text, page_index + 1, topic, topic))
                            body = []
                        topic = line[:100]
                    else:
                        body.append(line)
                if body:
                    text = _clean_text("\n".join(body))
                    if text:
                        sections.append(ExtractedSection(text, page_index + 1, topic, topic))
                char_count += sum(len(item.text) for item in sections if item.page_number == page_index + 1)
                if char_count > settings.max_extracted_characters:
                    raise ValueError("Extracted text exceeds the configured document text limit.")
    elif extension == ".docx" and content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if "word/document.xml" not in archive.namelist():
                    raise ValueError("The DOCX file is missing its document body.")
            document = Document(io.BytesIO(content))
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("The DOCX is corrupted or unreadable.") from exc
        topic: str | None = None
        pieces: list[str] = []
        for paragraph in document.paragraphs:
            text = _clean_text(paragraph.text)
            if not text:
                continue
            if paragraph.style and paragraph.style.name.lower().startswith("heading"):
                if pieces:
                    sections.append(ExtractedSection("\n".join(pieces), None, topic, topic))
                    pieces = []
                topic = text[:100]
            else:
                pieces.append(text)
        for table in document.tables:
            for row in table.rows:
                row_text = " | ".join(_clean_text(cell.text) for cell in row.cells)
                if row_text.strip(" |"):
                    pieces.append(row_text)
        if pieces:
            sections.append(ExtractedSection("\n".join(pieces), None, topic, topic))
    elif extension == ".txt" and content_type == "text/plain":
        try:
            text = _clean_text(content.decode("utf-8-sig"))
        except UnicodeDecodeError as exc:
            raise ValueError("Text documents must be valid UTF-8.") from exc
        topic: str | None = None
        body: list[str] = []
        for line in text.splitlines():
            line = line.strip()
            if line.startswith("#"):
                if body:
                    sections.append(ExtractedSection("\n".join(body), None, topic, topic))
                    body = []
                topic = line.lstrip("# ")[:100]
            elif line:
                body.append(line)
        if body:
            sections.append(ExtractedSection("\n".join(body), None, topic, topic))
    else:
        raise ValueError("Document processing currently supports PDF, DOCX, and TXT files.")

    total_text = sum(len(section.text) for section in sections)
    if total_text == 0:
        raise ValueError("No selectable text was found. This may be a scanned PDF; OCR is not enabled yet.")
    if total_text > settings.max_extracted_characters:
        raise ValueError("Extracted text exceeds the configured document text limit.")
    return sections, page_count


def split_sections(sections: list[ExtractedSection]) -> list[dict]:
    chunks: list[dict] = []
    for section in sections:
        paragraphs = [p.strip() for p in re.split(r"\n\s*\n|\n", section.text) if p.strip()]
        current = ""
        for paragraph in paragraphs:
            if len(paragraph) > _CHUNK_CHARS:
                if current:
                    chunks.append({"content": current, "page_number": section.page_number, "section_ref": section.section_ref, "topic": section.topic, "subtopic": section.subtopic})
                    current = ""
                for offset in range(0, len(paragraph), _CHUNK_CHARS - 150):
                    piece = paragraph[offset:offset + _CHUNK_CHARS]
                    if piece:
                        chunks.append({"content": piece, "page_number": section.page_number, "section_ref": section.section_ref, "topic": section.topic, "subtopic": section.subtopic})
                continue
            if current and len(current) + len(paragraph) + 1 > _CHUNK_CHARS:
                chunks.append({"content": current, "page_number": section.page_number, "section_ref": section.section_ref, "topic": section.topic, "subtopic": section.subtopic})
                current = ""
            current = f"{current}\n{paragraph}".strip()
        if current:
            chunks.append({"content": current, "page_number": section.page_number, "section_ref": section.section_ref, "topic": section.topic, "subtopic": section.subtopic})
    return chunks


def process_resource(resource_id: str) -> dict:
    """Fetch, extract, and index one resource. Repeated calls are content-hash cached."""
    if not _lock_for(resource_id).acquire(blocking=False):
        return {"resource_id": resource_id, "processing_status": "processing", "message": "Processing is already running."}
    client = None
    try:
        client = get_supabase_client()
        settings = get_settings()
        record = client.table("resources").select("id,title,category,course_id,storage_path,original_file_name,content_type,content_hash,processing_version,processing_status").eq("id", resource_id).limit(1).execute().data
        if not record:
            raise HTTPException(status_code=404, detail="Resource not found.")
        resource = record[0]
        if not resource.get("storage_path"):
            raise HTTPException(status_code=422, detail="This resource is a URL and has no stored document to process.")
        claimed = client.rpc("claim_resource_processing", {"target_resource_id": resource_id}).execute().data
        if not claimed:
            return {"resource_id": resource_id, "processing_status": "processing", "message": "Processing is already running."}
        file_response = client.storage.from_(settings.supabase_storage_bucket).download(resource["storage_path"])
        content = bytes(file_response)
        digest = hashlib.sha256(content).hexdigest()
        if resource.get("content_hash") == digest and resource.get("processing_version") == settings.ingestion_pipeline_version and resource.get("processing_status") == "completed":
            count = client.table("resource_chunks").select("id", count="exact").eq("resource_id", resource_id).execute().count or 0
            client.table("resources").update({"processing_status": "completed", "processing_error": None}).eq("id", resource_id).execute()
            return {"resource_id": resource_id, "processing_status": "completed", "content_hash": digest, "chunk_count": count, "cached": True}
        sections, _ = extract_document(content, resource.get("original_file_name") or resource["storage_path"], resource.get("content_type") or "")
        chunks = split_sections(sections)
        if not chunks:
            raise ValueError("No useful text passages could be created from this document.")
        # Upsert first to preserve the last good index if a transient request fails.
        payload = []
        for index, chunk in enumerate(chunks):
            payload.append({"resource_id": resource_id, "chunk_index": index, "page_number": chunk["page_number"], "content": chunk["content"], "content_hash": digest, "processing_version": settings.ingestion_pipeline_version, "topic": chunk["topic"], "subtopic": chunk["subtopic"], "section_ref": chunk["section_ref"], "metadata": {"title": resource["title"], "category": resource["category"], "course_id": resource["course_id"]}})
        client.table("resource_chunks").upsert(payload, on_conflict="resource_id,chunk_index").execute()
        client.table("resource_chunks").delete().eq("resource_id", resource_id).gte("chunk_index", len(payload)).execute()
        client.table("resources").update({"content_hash": digest, "processing_version": settings.ingestion_pipeline_version, "processing_status": "completed", "processing_error": None, "processed_at": datetime.now(timezone.utc).isoformat()}).eq("id", resource_id).execute()
        return {"resource_id": resource_id, "processing_status": "completed", "content_hash": digest, "chunk_count": len(chunks), "cached": False}
    except HTTPException:
        raise
    except Exception as exc:
        logger.error("Document processing failed for resource %s (%s)", resource_id, type(exc).__name__)
        if client is not None:
            try:
                safe_error = str(exc)[:500] if isinstance(exc, ValueError) else "Document processing failed. Retry the job or check server logs."
                client.table("resources").update({"processing_status": "failed", "processing_error": safe_error}).eq("id", resource_id).execute()
            except Exception as update_exc:
                logger.error("Could not persist processing failure for resource %s (%s)", resource_id, type(update_exc).__name__)
        raise
    finally:
        _lock_for(resource_id).release()
