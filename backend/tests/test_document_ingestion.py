import io
import zipfile

import pymupdf as fitz
import pytest
from docx import Document

from app.core.config import get_settings
from app.services.document_ingestion import ExtractedSection, extract_document, split_sections


def make_pdf(pages: list[str]) -> bytes:
    doc = fitz.open()
    for text in pages:
        page = doc.new_page()
        page.insert_text((72, 72), text)
    result = doc.tobytes()
    doc.close()
    return result


def test_pdf_extraction_preserves_page_references(monkeypatch):
    monkeypatch.setenv("MAX_DOCUMENT_PAGES", "10")
    sections, pages = extract_document(make_pdf(["FARADAY LAW\nMagnetic flux changes induce voltage.", "LENZ LAW\nThe induced current opposes change."]), "notes.pdf", "application/pdf")
    assert pages == 2
    assert any(item.page_number == 1 and item.topic == "FARADAY LAW" for item in sections)
    assert any(item.page_number == 2 and "opposes change" in item.text for item in sections)


def test_scanned_or_empty_pdf_is_not_success():
    with pytest.raises(ValueError, match="scanned PDF"):
        extract_document(make_pdf([""]), "scan.pdf", "application/pdf")


def test_corrupt_pdf_is_rejected():
    with pytest.raises(ValueError, match="corrupted"):
        extract_document(b"%PDF-invalid", "bad.pdf", "application/pdf")


def test_pdf_page_limit_is_enforced(monkeypatch):
    monkeypatch.setattr(get_settings(), "max_document_pages", 1)
    with pytest.raises(ValueError, match="1-page limit"):
        extract_document(make_pdf(["Page one", "Page two"]), "long.pdf", "application/pdf")


def test_docx_heading_structure_and_txt_utf8():
    document = Document()
    document.add_heading("Unit 1: Electromagnetism", level=1)
    document.add_paragraph("Faraday law relates flux change and induced EMF.")
    stream = io.BytesIO()
    document.save(stream)
    sections, _ = extract_document(stream.getvalue(), "notes.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert sections[0].topic == "Unit 1: Electromagnetism"
    text_sections, _ = extract_document("# Unit 2\n\nMaxwell equations".encode(), "notes.txt", "text/plain")
    assert text_sections[0].topic == "Unit 2"
    with pytest.raises(ValueError, match="UTF-8"):
        extract_document(b"\xff", "bad.txt", "text/plain")


def test_unsupported_extension_and_bad_docx_are_rejected():
    with pytest.raises(ValueError, match="supports PDF"):
        extract_document(b"image", "photo.png", "image/png")
    with pytest.raises(ValueError, match="corrupted"):
        extract_document(b"not a zip", "notes.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")


def test_chunk_splitting_keeps_page_and_section_and_bounds_passages():
    sections = [ExtractedSection("alpha " * 500, 3, "Chapter 1", "Chapter 1")]
    chunks = split_sections(sections)
    assert len(chunks) > 1
    assert all(len(chunk["content"]) <= 1200 for chunk in chunks)
    assert all(chunk["page_number"] == 3 and chunk["section_ref"] == "Chapter 1" for chunk in chunks)
